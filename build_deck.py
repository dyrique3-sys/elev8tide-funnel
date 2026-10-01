#!/usr/bin/env python3
"""ELEV8TIDE — The Growth System. Image-forward Keynote-compatible deck."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

A = "assets/"
TEAL      = RGBColor(0x0E,0x4A,0x43)
TEAL_DEEP = RGBColor(0x0A,0x33,0x2E)
GOLD      = RGBColor(0xC6,0xA1,0x5B)
GOLD_SOFT = RGBColor(0xD9,0xBC,0x86)
CREAM     = RGBColor(0xF6,0xF1,0xE7)
CREAM2    = RGBColor(0xFB,0xF8,0xF1)
INK       = RGBColor(0x12,0x2A,0x26)
WHITE     = RGBColor(0xFF,0xFF,0xFF)
MUTED     = RGBColor(0x5c,0x6f,0x6a)

DISP = "Didot"          # native macOS luxury serif
BODY = "Avenir Next"    # native macOS clean sans

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
blank = prs.slide_layouts[6]

def slide():
    return prs.slides.add_slide(blank)

def bg(s, color):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color

def set_alpha(fill_elem, pct):
    """pct = opacity 0-100 of the fill color."""
    srgb = fill_elem.find(qn('a:srgbClr'))
    if srgb is None: return
    a = srgb.makeelement(qn('a:alpha'), {'val': str(int(pct*1000))})
    srgb.append(a)

def full_image(s, path):
    # cover-fit image across whole slide (crop to fill)
    from PIL import Image
    iw, ih = Image.open(path).size
    sar = SW/SH; iar = iw/ih
    if iar > sar:  # image wider -> fit height, crop width
        h = SH; w = Emu(int(SH*iar))
        left = Emu(int((SW-w)/2)); top = Emu(0)
    else:
        w = SW; h = Emu(int(SW/iar))
        left = Emu(0); top = Emu(int((SH-h)/2))
    s.shapes.add_picture(path, left, top, width=w, height=h)

def overlay(s, color, opacity, left=0, top=0, w=None, h=None):
    w = w if w is not None else SW
    h = h if h is not None else SH
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(left)), Emu(int(top)), Emu(int(w)), Emu(int(h)))
    r.line.fill.background()
    r.fill.solid(); r.fill.fore_color.rgb = color
    set_alpha(r.fill.fore_color._xFill, opacity)
    r.shadow.inherit = False
    return r

def grad_overlay(s, color, op_left, op_right):
    """horizontal gradient overlay color op_left->op_right (left to right)."""
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0,0,SW,SH)
    r.line.fill.background(); r.shadow.inherit=False
    sp = r.fill._xPr
    for tag in ('a:noFill','a:solidFill','a:gradFill','a:blipFill','a:pattFill','a:grpFill'):
        e = sp.find(qn(tag))
        if e is not None: sp.remove(e)
    g = sp.makeelement(qn('a:gradFill'), {})
    lst = g.makeelement(qn('a:gsLst'), {})
    def gs(pos, op):
        s_ = g.makeelement(qn('a:gs'), {'pos':str(pos)})
        c = g.makeelement(qn('a:srgbClr'), {'val':'%02X%02X%02X'%(color[0],color[1],color[2])})
        a = c.makeelement(qn('a:alpha'), {'val':str(int(op*1000))})
        c.append(a); s_.append(c); return s_
    lst.append(gs(0, op_left)); lst.append(gs(100000, op_right))
    g.append(lst)
    lin = g.makeelement(qn('a:lin'), {'ang':'0','scaled':'1'})
    g.append(lin)
    # insert gradFill before line props
    ln = sp.find(qn('a:ln'))
    if ln is not None: sp.insert(list(sp).index(ln), g)
    else: sp.append(g)
    return r

def text(s, txt, left, top, w, h, size, color, font=BODY, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None, italic=False, line=1.08):
    tb = s.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = txt.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        r = p.add_run(); r.text = ln
        f = r.font
        f.size = Pt(size); f.name = font; f.bold = bold; f.italic = italic
        f.color.rgb = color
        if spacing is not None:
            rPr = r._r.get_or_add_rPr(); rPr.set('spc', str(spacing))
    return tb

def eyebrow(s, txt, left, top, w, color=GOLD, align=PP_ALIGN.LEFT):
    return text(s, txt.upper(), left, top, w, 0.4, 12.5, color, font=BODY, bold=True,
                align=align, spacing=380)

def goldline(s, left, top, w=0.64, color=GOLD):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(w), Pt(2.4))
    r.line.fill.background(); r.shadow.inherit=False
    r.fill.solid(); r.fill.fore_color.rgb = color
    return r

def logo(s, left, top, color=WHITE, size=22, gold=GOLD):
    tb = s.shapes.add_textbox(Inches(left), Inches(top), Inches(5), Inches(0.5))
    tf = tb.text_frame; tf.word_wrap=False
    p = tf.paragraphs[0]
    for ch, cl in [("ELEV", color),("8", gold),("TIDE", color)]:
        r = p.add_run(); r.text = ch
        r.font.size=Pt(size); r.font.name=DISP; r.font.bold=True; r.font.color.rgb=cl
        rPr=r._r.get_or_add_rPr(); rPr.set('spc','300')
    return tb

def pagenum(s, n, color=GOLD):
    text(s, f"{n:02d}", 12.4, 6.95, 0.8, 0.4, 10, color, font=BODY, spacing=200, align=PP_ALIGN.RIGHT)

# ---------------------------------------------------------------- 01 COVER
s = slide(); full_image(s, A+"life1.jpg")
grad_overlay(s, (0x0A,0x33,0x2E), 82, 58)
overlay(s, TEAL_DEEP, 30)  # unify
logo(s, 0.85, 0.7, color=WHITE, size=22)
eyebrow(s, "Advanced Wellness & Performance", 0.9, 2.7, 9)
text(s, "The Growth\nSystem.", 0.85, 3.05, 10, 2.4, 74, WHITE, font=DISP, line=0.95)
text(s, "How a luxury wellness brand becomes an\nautomated client-acquisition machine.", 0.9, 5.3, 9, 1.2, 19, GOLD_SOFT, font=BODY, line=1.2)
goldline(s, 0.9, 2.55, 0.9)

# ---------------------------------------------------------------- 02 VISION
s = slide(); full_image(s, A+"life4.jpg")
grad_overlay(s, (0x0A,0x33,0x2E), 30, 88)
eyebrow(s, "The Opportunity", 6.3, 2.3, 6, align=PP_ALIGN.LEFT)
text(s, "You built\nthe brand.", 6.25, 2.65, 6.5, 2, 52, WHITE, font=DISP, line=0.98)
text(s, "Now let's build\nthe machine.", 6.25, 4.45, 6.5, 1.6, 40, GOLD_SOFT, font=DISP, italic=True, line=1.0)
pagenum(s, 2)

# ---------------------------------------------------------------- 03 SYSTEM OVERVIEW (cream, 5 phases)
s = slide(); bg(s, CREAM2)
eyebrow(s, "The Build", 0.9, 0.75, 6, color=GOLD)
text(s, "Five phases. One engine.", 0.85, 1.1, 11, 1, 40, TEAL, font=DISP)
goldline(s, 0.9, 2.05, 0.7)
phases = [
    ("1","The Funnel","A luxury page that captures every lead."),
    ("2","Email + SMS","Automated nurture that warms leads to buyers."),
    ("3","Paid Ads","Targeted traffic from affluent buyers."),
    ("4","AI Agents","A 24/7 concierge that answers & books."),
    ("5","Automation","The whole engine, running itself."),
]
cw = 2.28; gap = 0.18; x0 = 0.9; y = 2.6
for i,(n,t,d) in enumerate(phases):
    x = x0 + i*(cw+gap)
    card = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(cw), Inches(3.5))
    card.fill.solid(); card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = GOLD; card.line.width = Pt(0.75); card.shadow.inherit=False
    text(s, n, x+0.05, y+0.35, cw-0.3, 1, 54, GOLD, font=DISP, align=PP_ALIGN.LEFT)
    text(s, t, x+0.28, y+1.55, cw-0.5, 0.6, 18, TEAL, font=DISP, bold=True)
    text(s, d, x+0.28, y+2.15, cw-0.5, 1.2, 11.5, MUTED, font=BODY, line=1.25)
pagenum(s, 3, color=GOLD)

# ---------------------------------------------------------------- Phase detail helper
def phase_slide(n, img, side, ey, big, big2, note, pg, tag):
    s = slide(); full_image(s, A+img)
    if side == "left":
        grad_overlay(s, (0x0A,0x33,0x2E), 90, 25)
        tx = 0.9
    else:
        grad_overlay(s, (0x0A,0x33,0x2E), 25, 90)
        tx = 6.5
    text(s, n, tx, 1.5, 4, 2, 120, GOLD, font=DISP, line=0.8)
    eyebrow(s, ey, tx+0.08, 3.65, 6.2)
    text(s, big+("\n"+big2 if big2 else ""), tx, 4.0, 6.4, 1.8, 40, WHITE, font=DISP, line=0.98)
    text(s, note, tx+0.03, big2 and 5.75 or 5.35, 6.0, 1.2, 14.5, GOLD_SOFT, font=BODY, line=1.25)
    # tag chip
    chip = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(tx+0.03), Inches(6.55), Inches(len(tag)*0.105+0.5), Inches(0.42))
    chip.fill.solid(); chip.fill.fore_color.rgb = GOLD; chip.line.fill.background(); chip.shadow.inherit=False
    text(s, tag.upper(), tx+0.03, 6.62, len(tag)*0.105+0.5, 0.3, 10, TEAL_DEEP, font=BODY, bold=True, align=PP_ALIGN.CENTER, spacing=200)
    pagenum(s, pg)

# 04 Phase 1
phase_slide("01","phone.jpg","right","Phase One · Live Now",
    "The luxury","front door.","A high-end landing page that captures every visitor's\nname, email & phone. Already built — sample is live.",4,"View the live sample")
# 05 Phase 2
phase_slide("02","life2.jpg","left","Phase Two",
    "Follow up,","automatically.","Welcome series + SMS that nurture every lead the\nmoment they join. 98% of texts get opened.",5,"Email + SMS")
# 06 Phase 3
phase_slide("03","life3.jpg","left","Phase Three",
    "Fuel the","fire.","Targeted Meta & Google ads that put ELEV8TIDE\nin front of the exact luxury buyer — at scale.",6,"Paid Ads")
# 07 Phase 4
phase_slide("04","life1.jpg","right","Phase Four",
    "A concierge","that never sleeps.","An AI team in the ELEV8TIDE voice — answering DMs,\nqualifying leads and booking calls 24/7.",7,"AI Agents")

# ---------------------------------------------------------------- 08 Phase 5 Automation (dark, diagram-ish)
s = slide(); bg(s, TEAL_DEEP)
overlay(s, TEAL, 18)
eyebrow(s, "Phase Five", 0.9, 0.8, 6)
text(s, "The complete engine.", 0.85, 1.15, 11, 1, 40, WHITE, font=DISP)
text(s, "Every piece connected — running itself.", 0.9, 2.05, 9, 0.6, 16, GOLD_SOFT, font=BODY, italic=True)
flow = ["Ads","Funnel","Nurture","Book","Follow-up"]
cw=2.1; gap=0.35; x0=0.95; y=3.3
for i,step in enumerate(flow):
    x = x0 + i*(cw+gap)
    c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(cw), Inches(1.5))
    c.fill.solid(); c.fill.fore_color.rgb = TEAL_DEEP
    c.line.color.rgb = GOLD; c.line.width=Pt(1); c.shadow.inherit=False
    text(s, step, x, y+0.5, cw, 0.6, 20, WHITE, font=DISP, bold=True, align=PP_ALIGN.CENTER)
    if i < len(flow)-1:
        text(s, "→", x+cw-0.02, y+0.42, 0.5, 0.6, 24, GOLD, font=BODY, align=PP_ALIGN.CENTER)
text(s, "One dashboard. Full visibility. A business that grows without adding hours.", 0.9, 5.6, 11, 0.8, 16, GOLD_SOFT, font=BODY, align=PP_ALIGN.LEFT)
pagenum(s, 8)

# ---------------------------------------------------------------- 09 CATALOG POSSIBILITY (flyer)
s = slide(); bg(s, CREAM2)
# flyer on right, text left
from PIL import Image as PImage
fw, fh = PImage.open(A+"flyer.jpg").size
dispH = Inches(7.5); dispW = Emu(int(7.5* (fw/fh) * 914400))
s.shapes.add_picture(A+"flyer.jpg", Emu(int(SW - dispW)), 0, height=dispH)
overlay(s, CREAM2, 100, left=0, top=0, w=Inches(6.2), h=SH)
eyebrow(s, "The Possibility", 0.9, 1.5, 6, color=GOLD)
text(s, "Your product line,\nelevated.", 0.85, 1.95, 6, 1.8, 36, TEAL, font=DISP, line=1.0)
text(s, "The same system can showcase the full\nELEV8TIDE catalog — premium, lab-tested,\nclinically focused — and turn interest\ninto orders automatically.", 0.9, 4.0, 5.5, 2, 15, MUTED, font=BODY, line=1.35)
goldline(s, 0.9, 3.75, 0.7)
pagenum(s, 9, color=GOLD)

# ---------------------------------------------------------------- 10 EXPERIENCE (spa)
s = slide(); full_image(s, A+"spa.jpg")
grad_overlay(s, (0x0A,0x33,0x2E), 88, 30)
eyebrow(s, "The Standard", 0.9, 2.5, 6)
text(s, "A brand that feels\nlike a private club.", 0.85, 2.9, 7.5, 2, 42, WHITE, font=DISP, line=1.0)
text(s, "Discretion, design and care in every touchpoint —\nonline exactly as it feels in person.", 0.9, 4.9, 7, 1, 15.5, GOLD_SOFT, font=BODY, line=1.25)
pagenum(s, 10)

# ---------------------------------------------------------------- 11 PRICING TIERS
s = slide(); bg(s, CREAM2)
eyebrow(s, "Investment", 0.9, 0.7, 6, color=GOLD)
text(s, "Choose your tide.", 0.85, 1.05, 11, 1, 40, TEAL, font=DISP)
text(s, "Start where it makes sense. Each tier unlocks — and pays for — the next.", 0.9, 1.95, 11, 0.5, 14.5, MUTED, font=BODY)
tiers = [
    ("LAUNCH","The Front Door","Funnel + lead capture",["Luxury 1-page funnel","Name / email / phone capture","Leads to your inbox","Mobile-optimized & live"], False),
    ("GROWTH","The Engine","Funnel + email/SMS + ads",["Everything in Launch","Email & SMS nurture","Meta + Google ads","Monthly optimization"], True),
    ("ELITE","Full Autopilot","Everything + AI agents",["Everything in Growth","24/7 AI concierge","End-to-end automation","Dedicated growth partner"], False),
]
cw=3.7; gap=0.4; x0=0.95; y=2.7
for i,(lbl,name,sub,feats,feat) in enumerate(tiers):
    x = x0+i*(cw+gap)
    card = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(cw), Inches(4.2))
    card.fill.solid(); card.fill.fore_color.rgb = TEAL if feat else WHITE
    card.line.color.rgb = GOLD; card.line.width = Pt(1 if feat else 0.75); card.shadow.inherit=False
    text(s, lbl, x+0.4, y+0.35, cw-0.8, 0.4, 12, GOLD, font=BODY, bold=True, spacing=300)
    text(s, name, x+0.4, y+0.75, cw-0.8, 0.6, 24, WHITE if feat else TEAL, font=DISP, bold=True)
    text(s, sub, x+0.4, y+1.35, cw-0.8, 0.4, 12, GOLD_SOFT if feat else MUTED, font=BODY, italic=True)
    for j,ft in enumerate(feats):
        text(s, "✦  "+ft, x+0.4, y+1.9+j*0.48, cw-0.8, 0.4, 12.5,
             (WHITE if feat else INK), font=BODY)
    if feat:
        rb = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x+cw/2-0.95), Inches(y-0.22), Inches(1.9), Inches(0.44))
        rb.fill.solid(); rb.fill.fore_color.rgb=GOLD; rb.line.fill.background(); rb.shadow.inherit=False
        text(s, "MOST POPULAR", x+cw/2-0.95, y-0.15, 1.9, 0.3, 9.5, TEAL_DEEP, font=BODY, bold=True, align=PP_ALIGN.CENTER, spacing=150)
text(s, "Exact pricing scoped to ELEV8TIDE's goals — final numbers in the proposal.", 0.9, 7.0, 11, 0.4, 11, MUTED, font=BODY, italic=True)

# ---------------------------------------------------------------- 12 CLOSE
s = slide(); full_image(s, A+"life3.jpg")
grad_overlay(s, (0x0A,0x33,0x2E), 55, 80)
overlay(s, TEAL_DEEP, 25)
logo(s, 0.9, 0.7, color=WHITE, size=20)
eyebrow(s, "Your Next Chapter Starts Here", 0.9, 3.0, 10, align=PP_ALIGN.LEFT)
text(s, "Elevate. Restore.\nOptimize.", 0.85, 3.4, 11, 2.2, 60, WHITE, font=DISP, line=0.98)
text(s, "Let's start with the front door — then scale from there.", 0.9, 5.7, 10, 0.6, 18, GOLD_SOFT, font=BODY, italic=True)
goldline(s, 0.9, 2.85, 0.9)

prs.save("ELEV8TIDE_Growth_System.pptx")
print("saved ELEV8TIDE_Growth_System.pptx —", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
